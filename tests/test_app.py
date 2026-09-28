import io
import os
import re
from datetime import date

import openpyxl
import pytest

from mtdapp import calc, fields as F, parser, report
from mtdapp.app import create_app
from mtdapp.db import DB


def filled_template(day, values_by_report):
    """Fill the downloadable 'all reports' template like a user would."""
    wb = openpyxl.load_workbook(io.BytesIO(report.template()))
    for ws in wb.worksheets:
        rep = next(k for k, r in F.REPORTS.items() if r["title"][:31] == ws.title)
        ws["B3"] = day
        labels = {f["label"]: f["key"] for f in F.fields(rep)}
        labels.update({f["label"] + " MTD": f["key"] + "_mtd" for f in F.fields(rep) if f.get("mtd")})
        for row in ws.iter_rows(min_row=7):
            key = labels.get(row[0].value)
            if key and key in values_by_report.get(rep, {}):
                row[1].value = values_by_report[rep][key]
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


SAMPLE = {
    F.PRODUCTION: dict(stoping_plan=1093, stoping_actual=1525, reef_dev_plan=9.6, reef_dev_actual=12.9,
                       waste_dev_plan=17.4, waste_dev_actual=26.4, trammed=7080, reef_hoisted=6385,
                       waste_hoisted=697, delivered=6510),
    F.ENGINEERING: dict(engineering_availability=0.8868, hoisting_availability=88.68, overall_belt_availability=1,
                        surface_belt_availability=100, eng_reef_hoisted=2298, eng_delivered=3843,
                        surface_stock=1123, ug_stock=1300, underlay=12.5, overlay=12.4, skips=0),
    F.LOGISTICS: dict(booked=44, full_down=41, empty_up=54, closing_ug=0, closing_surface=0,
                      explosives_down=0, vent_pipes_down=12),
}


def test_template_round_trip_classifies_and_reads_all_three_sheets():
    data = filled_template(date(2026, 9, 24), SAMPLE)
    recs, msgs = parser.parse_file(data, "daily.xlsx")
    assert not msgs
    by = {r["report"]: r for r in recs}
    assert set(by) == set(F.REPORTS)
    assert all(r["date"] == date(2026, 9, 24) for r in recs)
    assert by[F.PRODUCTION]["values"]["reef_hoisted"] == 6385
    assert by[F.ENGINEERING]["values"]["engineering_availability"] == pytest.approx(88.68)  # fraction -> %
    assert by[F.ENGINEERING]["values"]["overall_belt_availability"] == 100
    assert by[F.LOGISTICS]["values"]["vent_pipes_down"] == 12
    assert "bogeys_slung" not in by[F.LOGISTICS]["values"]  # blank stays "not reported"


def test_label_value_sheet_with_other_wording_and_units_in_text():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Material Car Report"
    ws.append(["Material Car Report Date", "24-Sep-2026"])
    ws.append(["Cars Down", "41"])
    ws.append(["Cars Up", 54])
    ws.append(["Left U/G", 0])
    ws.append(["Left Surface", "0 cars"])
    buf = io.BytesIO()
    wb.save(buf)
    recs, _ = parser.parse_file(buf.getvalue(), "cars.xlsx")
    assert recs[0]["report"] == F.LOGISTICS
    assert recs[0]["date"] == date(2026, 9, 24)
    assert recs[0]["values"] == {"full_down": 41, "empty_up": 54, "closing_ug": 0, "closing_surface": 0}


def test_table_layout_gives_one_record_per_row_and_date_from_filename():
    csv_data = b"Date,Reef Hoisted,U/G Trammed,Concentrator Delivered\n01-Sep-2026,2987,3826,3826\n02-Sep-2026,3946,4146,3837\n"
    recs, _ = parser.parse_file(csv_data, "production_backfill.csv")
    assert [r["date"] for r in recs] == [date(2026, 9, 1), date(2026, 9, 2)]
    assert recs[1]["values"]["reef_hoisted"] == 3946
    assert parser.to_date("Daily_Production_2026-09-23.xlsx") == date(2026, 9, 23)


def test_mtd_accumulates_and_reported_mtd_wins():
    db = DB(":memory:")
    db.save_day(F.PRODUCTION, "2026-09-01", {"reef_hoisted": 100, "stoping_actual": 50})
    db.save_day(F.PRODUCTION, "2026-09-02", {"reef_hoisted": 200, "stoping_actual_mtd": 10})  # series reset
    db.save_day(F.PRODUCTION, "2026-09-03", {"reef_hoisted": 300, "stoping_actual": 5})
    db.save_day(F.PRODUCTION, "2026-10-01", {"reef_hoisted": 7})
    rows = calc.series(db, F.PRODUCTION, "2026-09")
    assert [r["mtd"]["reef_hoisted"] for r in rows] == [100, 300, 600]
    assert [r["mtd"]["stoping_actual"] for r in rows] == [50, 10, 15]
    assert rows[-1]["mtd"]["delivered"] is None
    assert calc.series(db, F.PRODUCTION, "2026-10")[0]["mtd"]["reef_hoisted"] == 7


def test_recovery_matches_workbook_formulas():
    db = DB(":memory:")
    # 19 production reports whose MTD ends at 81,872 t, as at 24 Sep 2026
    for i, d in enumerate([1, 2, 3, 4, 5, 7, 8, 9, 10, 11, 12, 14, 15, 16, 17, 18, 21, 22, 24]):
        db.save_day(F.PRODUCTION, f"2026-09-{d:02d}", {"reef_hoisted": 4000} if i < 18 else {"reef_hoisted_mtd": 81872})
    db.save_day(F.ENGINEERING, "2026-09-24", {"eng_reef_hoisted": 2298, "underlay": 12.5, "overlay": 12.4})
    db.set_calendar("2026-09-24", 0, "PUBLIC HOLIDAY", None)
    d = calc.dashboard(db, "2026-09")
    r = d["recovery"]
    assert d["plan_to_date"] == 5218 * 19
    assert r["remaining_shifts"] == 15 and r["eq_days"] == 5
    assert r["required"] == pytest.approx(8672)
    assert r["projection"] == pytest.approx(103417, abs=1)
    assert d["call"]["hoisted_var"] == -2920
    assert d["spillage"]["inventory"] == pytest.approx(2120.8, abs=0.1)
    assert d["exceptions"][0]["name"] == "Reef Hoisted" and d["exceptions"][0]["status"] == "red"


def test_default_roster_september_2026():
    db = DB(":memory:")
    cal = {c["date"]: c for c in calc.roster(db, "2026-09")}
    assert cal["2026-09-04"]["shifts"] == 2          # Friday before OFF Saturday
    assert cal["2026-09-05"]["shifts"] == 0          # OFF Saturday
    assert cal["2026-09-11"]["shifts"] == 3          # Friday before working Saturday
    assert cal["2026-09-12"]["shifts"] == 2          # working Saturday
    assert cal["2026-09-13"]["shifts"] == 1          # Sunday night shift
    assert sum(c["shifts"] for d, c in cal.items() if d > "2026-09-24") == 15


@pytest.fixture
def client(tmp_path):
    app = create_app(data_dir=str(tmp_path))
    app.testing = True
    return app.test_client()


def test_upload_review_confirm_dashboard_export(client):
    data = filled_template(date(2026, 9, 24), SAMPLE)
    r = client.post("/upload", data={"files": (io.BytesIO(data), "daily.xlsx")}, content_type="multipart/form-data")
    assert r.status_code == 200 and b"Check the figures" in r.data
    token = re.search(rb'name="token" value="([0-9a-f]+)"', r.data).group(1).decode()
    form = {"token": token}
    for i in range(3):
        form[f"use_{i}"] = "on"
    r = client.post("/upload/confirm", data=form)
    assert r.status_code == 302
    r = client.get("/?month=2026-09")
    assert r.status_code == 200
    assert b"6,385 t" in r.data and b"-2,920 t" in r.data
    for page in ("/history?report=engineering&month=2026-09", "/calendar?month=2026-09", "/actions",
                 "/settings?month=2026-09", "/upload", "/day/production/2026-09-24", "/day/logistics"):
        assert client.get(page).status_code == 200, page
    x = client.get("/export?month=2026-09")
    wb = openpyxl.load_workbook(io.BytesIO(x.data))
    assert "MTD Dashboard" in wb.sheetnames and "Hoisting Calendar" in wb.sheetnames
    csv_r = client.get("/history?report=production&month=2026-09&format=csv")
    assert b"6385" in csv_r.data


def test_manual_edit_and_delete(client):
    r = client.post("/day/logistics", data={"date": "2026-09-25", "booked": "50", "full_down": "", "note": "hand"})
    assert r.status_code == 302
    db = client.application.db
    assert db.day(F.LOGISTICS, "2026-09-25")["values"] == {"booked": 50}
    client.post("/day/logistics/2026-09-25", data={"delete": "1"})
    assert db.day(F.LOGISTICS, "2026-09-25") is None


SAMPLE_WORKBOOK = os.environ.get("MTD_SAMPLE_WORKBOOK")


@pytest.mark.skipif(not SAMPLE_WORKBOOK, reason="set MTD_SAMPLE_WORKBOOK to the old consolidated workbook")
def test_history_import_reproduces_consolidated_workbook():
    from mtdapp import history_import
    db = DB(":memory:")
    history_import.import_consolidated(db, open(SAMPLE_WORKBOOK, "rb").read())
    d = calc.dashboard(db, "2026-09")
    assert d["as_at"] == "2026-09-24"
    flow = {f["key"]: f for f in d["flow"]}
    assert round(flow["reef_hoisted"]["pct"], 1) == 82.6
    assert round(flow["delivered"]["pct"], 1) == 84.1
    assert [round(m["pct"], 1) for m in d["mining"]] == [90.8, 98.6, 83.3]
    assert d["recovery"]["required"] == pytest.approx(8672)
    assert round(d["recovery"]["projection"]) == 103417
    assert d["logistics"]["mtd"]["full_down"] == 994 and d["logistics"]["mtd"]["empty_up"] == 1198


# ---- the real Thembelani report layouts -------------------------------------------------
def _xlsx(build):
    wb = openpyxl.Workbook()
    build(wb.active)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def engineering_sheet(ws):
    """Layout of the engineering daily snapshot (date left of the 'Date' label)."""
    from datetime import datetime
    rows = [(datetime(2026, 9, 18), "Date"), ("Eng Avail", 0.5636), ("Hoisting Avail", "55,91%"), (None, None),
            ("Overall Belt Avail", 1), ("Sfc Belts Avail", "100,00%"), (None, None),
            ("Hoist Today (dry)", 3513), ("Hoist MTD", "65 787"), (None, None),
            ("Mill Today", 2576), ("Mill MTD", 66524), (None, None),
            ("Surface Stocks", 1662), ("U/G Stocks", 1000), (None, None),
            ("Spillage", None), ("Underlay (meter)", "13,2"), ("Overlay (meter)", 12.8), ("Skips", 0)]
    for r in rows:
        ws.append(list(r))


def car_report_sheet(ws):
    """Layout of the Thembelane Shaft Car Report."""
    def put(ref, v):
        ws[ref] = v
    put("A1", "Date"); put("B1", "20/09/2026"); put("D1", "Thembelane Shaft Car Report")
    put("C2", "Empty Cars"); put("E2", "Full Cars")
    for col, h in zip("ABCDEFGH", ["LEVEL", "BOOKED", "Empties Up", "Empties Left\nUnderground", "Full Cars Down",
                                   "Full Cars Left on\nSurface", "Total Cars\nnot Down", "Remarks"]):
        put(f"{col}3", h)
    levels = [("14LEV", 5, 1, 0, 5, 0), ("15LEV", 6, 0, 0, 6, 0), ("16LEV", 3, 13, 0, 2, 0),
              ("17LEV", 7, 0, 0, 7, 0), ("18LEV", 3, 0, 0, 3, 0), ("19LEV", 29, 33, 18, 0, 29)]
    for i, row in enumerate(levels, 4):
        for col, v in zip("ABCDEF", row):
            put(f"{col}{i}", v)
    put("G6", 1); put("H6", "Car no:SC011 not found at bank area")
    put("H9", "29 Full cars left on surface and 18 empty cars left at the station due to 4 bogey's")
    for col, v in zip("ABCDEF", ["Total", 53, 47, 18, 23, 29]):
        put(f"{col}10", v)
    put("A11", "Total Full and Empty)"); put("C11", 70)
    put("A12", "Other Material")
    put("A13", "Sling work done"); put("D13", "Explosives Cars"); put("I13", "Vent Pipes"); put("M13", "Major Delays")
    for col, h in zip("ABCDEFIJKL", ["LEVEL", "Bogeys Down", "Bogey ID (number)", "Level", "Booked", "Down",
                                      "Level", "Booked", "Size (mm)", "Down"]):
        put(f"{col}14", h)
    put("M14", "17:10-Stop doing material cars doing all level report by Oasis banksman")
    for i, bid in enumerate(["BCP003 +BCP004", "BCP001 +BCP002", "F5+F6", "F3+F4"], 15):
        put(f"A{i}", 19); put(f"B{i}", 1); put(f"C{i}", bid)
    for i, lvl in enumerate(["14 Level", "15 Level", "16 Level", "17 Level", "18 Level", "19 Level"], 15):
        put(f"D{i}", lvl)
    put("A24", "Total"); put("B24", 4); put("D24", "Total"); put("E24", 0); put("F24", 0); put("J24", 0); put("L24", 0)


def test_engineering_snapshot_layout():
    recs, msgs = parser.parse_file(_xlsx(engineering_sheet), "Eng snapshot.xlsx")
    assert not msgs and len(recs) == 1
    r = recs[0]
    assert r["report"] == F.ENGINEERING and r["date"] == date(2026, 9, 18)
    assert r["values"] == {
        "engineering_availability": pytest.approx(56.36), "hoisting_availability": pytest.approx(55.91),
        "overall_belt_availability": 100, "surface_belt_availability": 100,
        "eng_reef_hoisted": 3513, "eng_hoist_mtd": 65787, "eng_delivered": 2576, "eng_mill_mtd": 66524,
        "surface_stock": 1662, "ug_stock": 1000, "underlay": 13.2, "overlay": 12.8, "skips": 0}


def test_shaft_car_report_layout_reads_total_rows():
    recs, msgs = parser.parse_file(_xlsx(car_report_sheet), "Car report.xlsx")
    assert not msgs and len(recs) == 1
    r = recs[0]
    assert r["report"] == F.LOGISTICS and r["date"] == date(2026, 9, 20)
    assert r["values"] == {"booked": 53, "empty_up": 47, "closing_ug": 18, "full_down": 23, "closing_surface": 29,
                           "explosives_down": 0, "vent_pipes_down": 0, "bogeys_slung": 4}
    assert "SC011" in r["note"] and "Oasis banksman" in r["note"]


def test_pdf_words_are_rebuilt_into_label_daily_mtd_rows():
    def w(text, x0, x1, top):
        return dict(text=text, x0=x0, x1=x1, top=top)
    words = [w("Thembelani", 184, 220, 60), w("Daily", 222, 240, 60), w("Report", 242, 265, 60),
             w("18", 267, 275, 60), w("Sept", 277, 292, 60), w("2026", 294, 310, 60),
             w("Month", 312, 330, 60), w("to", 332, 338, 60), w("Date", 340, 355, 60),
             w("Actual", 184.9, 204.6, 409), w("Reef", 206, 220, 409.4), w("Tonnes", 222, 245, 409),
             w("3", 343.9, 347.3, 409), w("512", 348.9, 359.0, 409), w("65", 404.2, 411.0, 409), w("427", 412.6, 422.7, 409),
             w("%", 184.9, 190, 420), w("Achieved", 192, 220, 420), w("67,3%", 343, 359, 420), w("78,4%", 404, 422, 420)]
    grid = parser.words_to_grid(words)
    assert grid[1] == ["Actual Reef Tonnes", "3512", "65427"]
    assert grid[2] == ["% Achieved", "67,3%", "78,4%"]
    recs = parser.parse_grid(F.PRODUCTION, grid)
    assert recs[0]["date"] == date(2026, 9, 18)
    assert recs[0]["values"] == {"reef_hoisted": 3512, "reef_hoisted_mtd": 65427}


def test_south_african_number_formats():
    assert parser.to_number("3 513") == 3513
    assert parser.to_number("56,36%") == pytest.approx(56.36)
    assert parser.to_number("-1 706") == -1706
    assert parser.to_number("1,123 t") == 1123
    assert parser.to_number("1 234,5") == 1234.5
    assert parser.to_number("—") is None


SAMPLE_PDF = os.environ.get("MTD_SAMPLE_PDF")


@pytest.mark.skipif(not SAMPLE_PDF, reason="set MTD_SAMPLE_PDF to a Daily Production Report PDF")
def test_real_production_pdf():
    recs, msgs = parser.parse_file(open(SAMPLE_PDF, "rb").read(), os.path.basename(SAMPLE_PDF))
    assert recs and recs[0]["report"] == F.PRODUCTION
    v = recs[0]["values"]
    assert {"reef_hoisted", "reef_hoisted_mtd", "trammed_mtd", "delivered_mtd", "stoping_actual_mtd"} <= set(v)


def _review(client, files):
    r = client.post("/upload", data={"files": files}, content_type="multipart/form-data")
    assert r.status_code == 200, r.data[:500]
    token = re.search(rb'name="token" value="([0-9a-f]+)"', r.data).group(1).decode()
    return r, token


def test_engineering_date_is_the_reporting_date_for_the_whole_upload(client):
    r, token = _review(client, [(io.BytesIO(_xlsx(car_report_sheet)), "Cars.xlsx"),
                                (io.BytesIO(_xlsx(engineering_sheet)), "Eng.xlsx")])
    # the car report says 20/09/2026, the engineering report 18/09/2026
    assert r.data.count(b'value="2026-09-18"') == 2 and b'value="2026-09-20"' not in r.data
    assert b"set to the engineering report date" in r.data
    client.post("/upload/confirm", data={"token": token, "use_0": "on", "use_1": "on",
                                         "date_0": "2026-09-18", "date_1": "2026-09-18"})
    db = client.application.db
    assert db.day(F.LOGISTICS, "2026-09-18")["values"]["booked"] == 53
    assert db.day(F.ENGINEERING, "2026-09-18")["values"]["eng_reef_hoisted"] == 3513


def test_without_engineering_report_different_dates_are_flagged(client):
    other = openpyxl.load_workbook(io.BytesIO(_xlsx(car_report_sheet)))
    other.active["B1"] = "21/09/2026"
    buf = io.BytesIO()
    other.save(buf)
    r, _ = _review(client, [(io.BytesIO(_xlsx(car_report_sheet)), "Cars.xlsx"),
                            (io.BytesIO(filled_template(date(2026, 9, 18), {F.PRODUCTION: SAMPLE[F.PRODUCTION]})),
                             "prod.xlsx")])
    assert b"No engineering report in this upload" in r.data


def test_values_corrected_on_the_review_screen_are_saved(client):
    r, token = _review(client, [(io.BytesIO(_xlsx(engineering_sheet)), "Eng.xlsx")])
    assert b'name="val_0_eng_reef_hoisted" value="3,513"' in r.data
    client.post("/upload/confirm", data={"token": token, "use_0": "on", "date_0": "2026-09-18",
                                         "val_0_eng_reef_hoisted": "3 600", "val_0_skips": "",
                                         "val_0_underlay": "13,2", "note_0": "checked"})
    rec = client.application.db.day(F.ENGINEERING, "2026-09-18")
    assert rec["values"] == {"eng_reef_hoisted": 3600, "underlay": 13.2}
    assert rec["note"] == "checked"


def test_unreadable_total_falls_back_to_sum_of_levels_and_mismatch_is_flagged():
    grid = [["LEVEL", "BOOKED", "Empties Up"],
            ["14LEV", "5", "1"], ["15LEV", "6", "0"],
            ["Total", None, "4"]]
    vals, warnings = parser.column_totals(F.LOGISTICS, grid)
    assert vals == {"booked": 11, "empty_up": 4}
    assert any("sum of the rows above" in w for w in warnings)
    assert any("Empty Cars Up: Total says 4 but the rows above add up to 1" in w for w in warnings)


def test_ocr_cleanup_snaps_labels_and_blanks_garbled_numbers():
    grid = [["Bogevs Down", "Full Cars Lefton Surface", "2a", "29", "56,36%", "18/09/2026 Date", "14LEV"]]
    cleaned = parser.clean_ocr_grid(grid)[0]
    assert parser.norm(cleaned[0]) == "bogeysdown"
    assert parser.norm(cleaned[1]) == "fullcarsleftonsurface"
    assert cleaned[2:] == [None, "29", "56,36%", "18/09/2026 Date", "14LEV"]


def _tesseract_available():
    try:
        from mtdapp import ocr
        ocr._libs()
        return True
    except Exception:
        return False


FONT = next((f for f in ("/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf",
                         "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf", "C:/Windows/Fonts/arialbd.ttf")
             if os.path.exists(f)), None)


@pytest.mark.skipif(not _tesseract_available() or not FONT, reason="Tesseract OCR or a TrueType font not available")
def test_picture_of_a_table_is_read():
    """A ruled table drawn like the engineering snapshot is read cell by cell."""
    from PIL import Image, ImageDraw, ImageFont
    rows = [("18/09/2026", "Date"), ("Eng Avail", "56,36%"), ("Hoist Today", "3 513"), ("Mill Today", "2 576"),
            ("Surface Stocks", "1 662"), ("Underlay (meter)", "13,2"), ("Skips", "0")]
    img = Image.new("RGB", (680, 60 * len(rows) + 20), "white")
    draw = ImageDraw.Draw(img)
    font = ImageFont.truetype(FONT, 32)
    for i, (a, b) in enumerate(rows):
        y = 10 + 60 * i
        draw.rectangle((10, y, 430, y + 60), outline="black", width=2)
        draw.rectangle((430, y, 670, y + 60), outline="black", width=2)
        draw.text((20, y + 12), a, font=font, fill="black")
        draw.text((660 - draw.textlength(b, font=font), y + 12), b, font=font, fill="black")
    buf = io.BytesIO()
    img.save(buf, "PNG")
    recs, msgs = parser.parse_file(buf.getvalue(), "eng.png")
    assert recs and recs[0]["report"] == F.ENGINEERING and recs[0]["from_image"]
    assert recs[0]["date"] == date(2026, 9, 18)
    assert recs[0]["values"] == {"engineering_availability": pytest.approx(56.36), "eng_reef_hoisted": 3513,
                                 "eng_delivered": 2576, "surface_stock": 1662, "underlay": 13.2, "skips": 0}


SAMPLE_IMAGES = os.environ.get("MTD_SAMPLE_IMAGES")   # "engineering.jpg,carreport.jpg"


@pytest.mark.skipif(not SAMPLE_IMAGES, reason="set MTD_SAMPLE_IMAGES=eng.jpg,cars.jpg to test real pictures")
def test_real_report_pictures():
    eng, cars = SAMPLE_IMAGES.split(",")
    e = parser.parse_file(open(eng, "rb").read(), os.path.basename(eng))[0][0]
    c = parser.parse_file(open(cars, "rb").read(), os.path.basename(cars))[0][0]
    assert e["report"] == F.ENGINEERING and c["report"] == F.LOGISTICS
    assert e["values"]["eng_reef_hoisted"] == 3513 and e["values"]["engineering_availability"] == pytest.approx(56.36)
    assert {k: c["values"][k] for k in ("booked", "empty_up", "closing_ug", "full_down", "closing_surface",
                                        "bogeys_slung")} == dict(booked=53, empty_up=47, closing_ug=18, full_down=23,
                                                                 closing_surface=29, bogeys_slung=4)


def test_password_protects_every_page_except_health_check(tmp_path, monkeypatch):
    monkeypatch.setenv("MTD_PASSWORD", "Shaft2026!")
    c = create_app(data_dir=str(tmp_path)).test_client()
    assert c.get("/healthz").data == b"ok"
    for page in ("/", "/upload", "/history", "/backup", "/export"):
        r = c.get(page)
        assert r.status_code == 302 and "/login" in r.headers["Location"], page
    assert b"Wrong password" in c.post("/login", data={"password": "nope"}).data
    r = c.post("/login?next=/history", data={"password": "Shaft2026!"})
    assert r.status_code == 302 and r.headers["Location"].endswith("/history")
    assert c.get("/").status_code == 200
    assert c.post("/login?next=//evil.example", data={"password": "Shaft2026!"}).headers["Location"] == "/"
    c.get("/logout")
    assert c.get("/").status_code == 302


def test_backup_and_restore(client):
    db = client.application.db
    db.save_day(F.LOGISTICS, "2026-09-18", {"booked": 53})
    backup = client.get("/backup").data
    assert backup.startswith(b"SQLite format 3")
    db.save_day(F.LOGISTICS, "2026-09-18", {"booked": 1})
    assert b"not an MTD backup" in client.post("/restore", data={"backup": (io.BytesIO(b"junk"), "x.db")},
                                                follow_redirects=True).data
    client.post("/restore", data={"backup": (io.BytesIO(backup), "mtd-backup.db")})
    assert db.day(F.LOGISTICS, "2026-09-18")["values"] == {"booked": 53}
