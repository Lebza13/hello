"""One-off import of the old consolidated MTD workbook, so the app starts with
the history already captured in the spreadsheet it replaces."""
import io
import re

import openpyxl

from . import calc, fields as F
from .parser import to_date, to_number


def is_consolidated(data):
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True)
    except Exception:
        return False
    names = set(wb.sheetnames)
    return {"MTD Trend", "Logistics MTD"} <= names


def _num_after(text, pattern):
    m = re.search(pattern, text or "")
    return to_number(m.group(1)) if m else None


def import_consolidated(db, data, filename="consolidated.xlsx"):
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    dash = wb["MTD Dashboard"]
    header = str(dash["A3"].value or "")
    m = re.search(r"(January|February|March|April|May|June|July|August|September|October|November|December)\s+(\d{4})",
                  header)
    year = int(m.group(2)) if m else None
    from datetime import date
    default = date(year, 1, 1) if year else None
    log = []

    # ---- production (MTD Trend) --------------------------------------------
    ws = wb["MTD Trend"]
    prod = {}
    prev = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        d = to_date(row[0], year) if row[0] and not str(row[0]).startswith("*") else None
        if not d or row[2] is None:
            continue
        vals = {"reef_hoisted_mtd": row[2], "trammed_mtd": row[3], "delivered_mtd": row[4],
                "stoping_plan_mtd": row[5], "stoping_actual_mtd": row[6]}
        for k in ("reef_hoisted", "trammed", "delivered", "stoping_plan", "stoping_actual"):
            cur, p = vals[k + "_mtd"], prev.get(k)
            if cur is not None:
                diff = cur - (p or 0)
                vals[k] = diff if diff >= 0 else None   # series resets are not daily figures
                prev[k] = cur
        note = "Combined production source (see original report)" if str(row[0]).endswith("*") else None
        prod[d.isoformat()] = (vals, note)
    latest_prod = max(prod) if prod else None

    # Dashboard text tiles give the latest mining + waste figures.
    if latest_prod:
        vals, note = prod[latest_prod]
        for cell, a, p in (("A", "stoping_actual", "stoping_plan"), ("E", "reef_dev_actual", "reef_dev_plan"),
                           ("I", "waste_dev_actual", "waste_dev_plan")):
            t7, t9 = str(dash[f"{cell}7"].value or ""), str(dash[f"{cell}9"].value or "")
            mtd = _num_after(t7, r"MTD\s+([\d,.]+)")
            pc = _num_after(t7, r"\|\s*([\d.]+)%")
            if mtd is not None:
                vals[a + "_mtd"] = mtd
                if pc and vals.get(p + "_mtd") is None:
                    vals[p + "_mtd"] = round(mtd / (pc / 100), 1)
            day = _num_after(t9, r"combined\s+([\d,.]+)")
            dplan = _num_after(t9, r"plan\s+([\d,.]+)")
            if day is not None:
                vals[a] = day
            if dplan is not None:
                vals[p] = dplan
        waste_mtd = _num_after(str(dash["H16"].value), r"MTD\s+([\d,.]+)")
        waste_pct = _num_after(str(dash["H16"].value), r"\|\s*([\d.]+)%")
        if waste_mtd is not None:
            vals["waste_hoisted_mtd"] = waste_mtd
            vals["waste_hoisted"] = _num_after(str(dash["H18"].value), r"waste\s+([\d,.]+)")
        for d, (v, nt) in prod.items():
            db.save_day(F.PRODUCTION, d, {k: x for k, x in v.items() if x is not None},
                        source_file=filename, note=nt)
        log.append(f"Production: {len(prod)} days imported")

    # ---- engineering (Spillage table + dashboard snapshot) ----------------------
    ws = wb["Spillage"]
    eng = {}
    for row in ws.iter_rows(min_row=2, max_row=40, values_only=True):
        d = to_date(row[0], year) if isinstance(row[0], str) and re.match(r"\d", row[0]) else None
        if not d:
            continue
        v = {"skips": to_number(row[1]), "underlay": to_number(row[4]), "overlay": to_number(row[5])}
        eng[d.isoformat()] = {k: x for k, x in v.items() if x is not None}
    snap_date = None
    snap_title = str(dash["A26"].value or "")
    sd = to_date(snap_title.split("|")[-1], year) if "|" in snap_title else None
    if sd:
        snap_date = sd.isoformat()
        v = eng.setdefault(snap_date, {})
        for label_cell, value_cell, key in (("A22", "A23", "engineering_availability"),
                                            ("D22", "D23", "hoisting_availability"),
                                            ("G22", "G23", "overall_belt_availability"),
                                            ("K22", "K23", "surface_belt_availability"),
                                            ("A28", "C28", "surface_stock"), ("A29", "C29", "ug_stock"),
                                            ("E28", "G28", "underlay"), ("E29", "G29", "overlay"),
                                            ("E30", "G30", "skips")):
            n = to_number(dash[value_cell].value)
            if n is not None:
                v[key] = n
        v["eng_reef_hoisted"] = _num_after(str(dash["D18"].value), r"([\d,]+)\s*t")
        v["eng_delivered"] = _num_after(str(dash["K18"].value), r"([\d,]+)\s*t")
    for d, v in eng.items():
        db.save_day(F.ENGINEERING, d, {k: x for k, x in v.items() if x is not None}, source_file=filename)
    log.append(f"Engineering: {len(eng)} days imported")

    # ---- logistics (Material Car Report table) ----------------------------------
    ws = wb["Logistics MTD"]
    n_log = 0
    last_log = None
    for row in ws.iter_rows(min_row=25, max_row=49, values_only=True):
        d = to_date(row[0], year) if isinstance(row[0], str) and re.match(r"\d", row[0]) else None
        if not d:
            continue
        v = {"booked": to_number(row[2]), "full_down": to_number(row[3]), "empty_up": to_number(row[4]),
             "closing_ug": to_number(row[6]), "closing_surface": to_number(row[7])}
        v = {k: x for k, x in v.items() if x is not None}
        if not any(k in v for k in ("booked", "full_down", "empty_up")):
            continue  # rows without a Material Car Report are left unimputed
        db.save_day(F.LOGISTICS, d.isoformat(), v, source_file=filename, note=row[8])
        n_log += 1
        last_log = d.isoformat()
    if last_log:
        rec = db.day(F.LOGISTICS, last_log)["values"]
        # The latest "MTD TOTAL" row includes days whose detail was never captured.
        for row in ws.iter_rows(min_row=25, max_row=49, values_only=True):
            if str(row[0] or "").startswith("MTD TOTAL"):
                for i, key in ((2, "booked"), (3, "full_down"), (4, "empty_up")):
                    if to_number(row[i]) is not None:
                        rec[key + "_mtd"] = to_number(row[i])
        for label, key in (("MTD explosives down", "explosives_down"), ("MTD vent pipes down", "vent_pipes_down"),
                           ("MTD bogeys slung", "bogeys_slung")):
            for row in ws.iter_rows(min_row=4, max_row=20, values_only=True):
                if row[0] == label and to_number(row[1]) is not None:
                    rec[key + "_mtd"] = to_number(row[1])
        for row in ws.iter_rows(min_row=50, max_row=62, values_only=True):
            if row[0] == "Explosives down":
                rec["explosives_down"] = to_number(row[1])
            if row[0] == "Vent pipes down":
                rec["vent_pipes_down"] = to_number(row[1])
        db.save_day(F.LOGISTICS, last_log, {k: x for k, x in rec.items() if x is not None})
    log.append(f"Material cars: {n_log} days imported")

    # ---- month settings + calendar -----------------------------------------------
    month = (latest_prod or snap_date or "")[:7]
    if "Hoisting Calendar" in wb.sheetnames and month:
        cws = wb["Hoisting Calendar"]
        for key, cell in (("monthly_reef_plan", "B5"), ("daily_call", "E5"), ("planned_hoisting_days", "H5")):
            n = to_number(cws[cell].value)
            if n:
                db.set_setting(month, key, n)
        if latest_prod and waste_mtd and waste_pct:
            db.set_setting(month, "waste_daily_call", round(waste_mtd / (waste_pct / 100) / len(prod), 2))
        default = {c["date"]: c for c in calc.roster(db, month)}
        n_cal = 0
        for row in cws.iter_rows(min_row=11, max_row=45, values_only=True):
            d = to_date(row[0])
            if not d or d.isoformat() not in default:
                continue
            base = default[d.isoformat()]
            shifts = to_number(row[4])
            if row[2] != base["status"] or (shifts is not None and shifts != base["shifts"]) or row[7]:
                db.set_calendar(d.isoformat(), shifts if shifts is not None else base["shifts"],
                                row[2] or base["status"], row[7])
                n_cal += 1
        log.append(f"Calendar: settings + {n_cal} day notes imported")

    # ---- management actions ------------------------------------------------------------
    if "Management Actions" in wb.sheetnames and not db.actions(include_closed=True):
        n_act = 0
        for row in wb["Management Actions"].iter_rows(min_row=2, values_only=True):
            if row[1]:
                icon = str(row[0] or "")
                st = "red" if "🔴" in icon else "amber" if "🟡" in icon else "green"
                db.save_action(dict(status=st, finding=row[1], action=row[2], owner=row[3], due=row[4],
                                    measure=row[5], closed=0))
                n_act += 1
        log.append(f"Management actions: {n_act} imported")
    return log
