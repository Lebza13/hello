"""Excel input templates and the consolidated month export."""
import io
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import calc, fields as F

HEAD = PatternFill("solid", fgColor="1F3A5F")
SUB = PatternFill("solid", fgColor="DCE6F1")
FILLS = {"green": "C6EFCE", "amber": "FFEB9C", "red": "FFC7CE"}
WHITE = Font(bold=True, color="FFFFFF")
BOLD = Font(bold=True)


def _bytes(wb):
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _widths(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _template_sheet(ws, report, day):
    spec = F.REPORTS[report]
    ws["A1"] = spec["title"]
    ws["A1"].font = Font(bold=True, size=14)
    ws["A3"], ws["B3"] = "Date", day
    ws["B3"].number_format = "dd-mmm-yyyy"
    ws["A4"] = "Note"
    ws["A3"].font = ws["A4"].font = BOLD
    ws["A6"], ws["B6"], ws["C6"] = "Figure", "Value", "Unit"
    for c in "ABC":
        ws[f"{c}6"].fill, ws[f"{c}6"].font = HEAD, WHITE
    r = 7
    for f in spec["fields"]:
        ws.cell(r, 1, f["label"])
        ws.cell(r, 3, f["unit"])
        r += 1
    mtd = [f for f in spec["fields"] if f.get("mtd")]
    if mtd:
        r += 1
        ws.cell(r, 1, "Optional: month-to-date as stated on the source report "
                      "(leave blank and the app adds up the daily figures)").font = Font(italic=True)
        r += 1
        for f in mtd:
            ws.cell(r, 1, f["label"] + " MTD")
            ws.cell(r, 3, f["unit"])
            r += 1
    _widths(ws, [44, 16, 10])


def template(report=None):
    """Blank input workbook. report=None gives one workbook with all three sheets."""
    wb = Workbook()
    wb.remove(wb.active)
    for rep in ([report] if report else list(F.REPORTS)):
        ws = wb.create_sheet(F.REPORTS[rep]["title"][:31])
        _template_sheet(ws, rep, date.today())
    return _bytes(wb)


def _table(ws, start_row, headers, rows, widths=None):
    for i, h in enumerate(headers, 1):
        c = ws.cell(start_row, i, h)
        c.fill, c.font = HEAD, WHITE
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for r, row in enumerate(rows, start_row + 1):
        for i, v in enumerate(row, 1):
            c = ws.cell(r, i, v)
            if isinstance(v, float):
                c.number_format = "#,##0.0" if abs(v) < 100 else "#,##0"
    if widths:
        _widths(ws, widths)
    return start_row + len(rows) + 1


def _status_cell(cell, st):
    if st in FILLS:
        cell.fill = PatternFill("solid", fgColor=FILLS[st])


def export_month(db, month):
    d = calc.dashboard(db, month)
    wb = Workbook()
    ws = wb.active
    ws.title = "MTD Dashboard"
    if not d:
        ws["A1"] = f"No data for {month}"
        return _bytes(wb)
    ws["A1"] = f"MTD PERFORMANCE DASHBOARD — {month}"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"Position as at {d['as_at']} | {d['production_days']} production reports this month"
    rows = []
    for m in d["mining"]:
        rows.append((m["name"], m["plan"], m["mtd"], m["pct"] / 100 if m["pct"] is not None else None,
                     m["unit"], m["status"]))
    for f in d["flow"]:
        rows.append((f["name"], f["plan"], f["mtd"], f["pct"] / 100 if f["pct"] is not None else None,
                     "t", f["status"]))
    end = _table(ws, 4, ["Indicator", "Plan MTD", "Actual MTD", "% of plan", "Unit", "Status"], rows,
                 [34, 14, 14, 12, 8, 10])
    for r in range(5, end):
        ws.cell(r, 4).number_format = "0.0%"
        _status_cell(ws.cell(r, 6), ws.cell(r, 6).value)
    rc = d["recovery"]
    rows = [("Month plan", rc["plan"], "t"), ("MTD reef hoisted", rc["mtd"], "t"), ("Balance", rc["balance"], "t"),
            ("Roster shifts remaining", rc["remaining_shifts"], "shifts"),
            ("Equivalent hoisting days left", rc["eq_days"], "days"),
            ("Required rate", rc["required"], "t/day"),
            ("Required uplift vs call", (rc["uplift"] or 0) / 100, "%"),
            ("Projection", rc["projection"], "t"), ("Forecast shortfall", rc["shortfall"], "t"),
            ("Daily call", d["call"]["call"], "t"),
            ("Reef hoisted (latest day)", d["call"]["hoisted_day"], "t"),
            ("Delivered (latest day)", d["call"]["delivered_day"], "t")]
    rows += [(n, v, "%") for n, v in d["availability"]]
    rows += [("Surface stock", d["stocks"]["surface"], "t"), ("U/G stock", d["stocks"]["ug"], "t"),
             ("Total stock", d["stocks"]["total"], "t")]
    end = _table(ws, end + 2, ["Recovery & engineering", "Value", "Unit"], rows)
    for r in range(1, end):
        if ws.cell(r, 3).value == "%" and ws.cell(r, 1).value == "Required uplift vs call":
            ws.cell(r, 2).number_format = "+0.0%"
    _table(ws, end + 2, ["Top exceptions", "Detail", "Status"],
           [(x["name"], x["detail"], x["status"]) for x in d["exceptions"]])

    # trend
    ws = wb.create_sheet("MTD Trend")
    _table(ws, 1, ["Date", "Reef Hoisted Plan", "Reef Hoisted Actual", "U/G Trammed", "Concentrator Delivered",
                   "Stoping Plan", "Stoping Actual", "Projection"],
           [(t["date"], t["plan"], t["hoisted"], t["trammed"], t["delivered"], t["stoping_plan"], t["stoping"],
             t["projection"]) for t in d["trend"]], [12] + [16] * 7)

    for rep, sheet in ((F.PRODUCTION, "Production Daily"), (F.ENGINEERING, "Engineering Daily"),
                       (F.LOGISTICS, "Logistics MTD")):
        ws = wb.create_sheet(sheet)
        fl = F.fields(rep)
        heads = ["Date"] + [f"{f['label']} ({f['unit']})" for f in fl] + \
                [f"{f['label']} MTD" for f in fl if f.get("mtd")] + ["Note"]
        body = []
        for r in calc.series(db, rep, month):
            body.append([r["date"]] + [r["values"].get(f["key"]) for f in fl] +
                        [r["mtd"].get(f["key"]) for f in fl if f.get("mtd")] + [r["note"]])
        _table(ws, 1, heads, body, [12] + [14] * (len(heads) - 2) + [40])
        ws.row_dimensions[1].height = 45

    sp = d["spillage"]
    ws = wb.create_sheet("Spillage")
    end = _table(ws, 1, ["Date", "Skips", "Tonnes Removed", "Cumulative Tonnes", "Underlay m", "Overlay m"],
                 [(r["date"], r["skips"], r["removed"], r["cumulative"], r["underlay"], r["overlay"])
                  for r in sp["rows"]], [14, 10, 14, 16, 12, 12])
    model = [("Gross geometric capacity", sp.get("gross"), "t"), ("Critical spillage capacity", sp.get("critical"), "t"),
             ("Current average clearance", sp.get("clearance"), "m"), ("Estimated spillage depth", sp.get("depth"), "m"),
             ("Estimated in-shaft inventory", sp.get("inventory"), "t"),
             ("Inventory % of critical", sp.get("pct_critical"), "%"),
             ("Margin to critical level", sp.get("margin"), "t"), ("MTD spillage removed", sp.get("removed"), "t"),
             ("Removal rate (calendar day)", sp.get("cal_rate"), "t/day"),
             ("Days to clear at this rate", sp.get("days_to_clear"), "days")]
    _table(ws, end + 2, ["Shaft-bottom model", "Value", "Unit"], model)

    ws = wb.create_sheet("Management Actions")
    _table(ws, 1, ["Status", "Finding", "Required action", "Owner", "Due", "Success measure"],
           [(a["status"], a["finding"], a["action"], a["owner"], a["due"], a["measure"]) for a in db.actions()],
           [10, 50, 50, 22, 12, 30])

    ws = wb.create_sheet("Hoisting Calendar")
    _table(ws, 1, ["Date", "Day", "Roster Status", "Shift Pattern", "Available Shifts", "Period", "Notes"],
           [(c["date"], c["day"], c["status"], c["pattern"], c["shifts"],
             "ELAPSED" if c["date"] <= d["as_at"] else "REMAINING", c["note"]) for c in d["calendar"]],
           [12, 12, 30, 28, 10, 12, 60])
    return _bytes(wb)
